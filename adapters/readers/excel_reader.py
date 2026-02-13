"""Excel (XLSX) reader implementation.

Reads agency-specific Excel files based on mapping definitions
and converts them to the CIR model.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

import openpyxl

from adapters.mapping import MappingDefinition
from adapters.readers.base import BaseReader
from core.models import DMP


class ExcelReader(BaseReader):
    """Reads data from an Excel file based on a mapping definition."""

    def __init__(self, mapping: MappingDefinition) -> None:
        """Initialize the ExcelReader.

        Args:
            mapping: Mapping definition linking cell coordinates to CIR fields.
        """
        self._mapping = mapping

    def read(self, input_path: Path) -> DMP:
        """Read an Excel file and return the CIR representation."""
        wb = openpyxl.load_workbook(input_path, data_only=True)
        raw: dict[str, str | None] = {}

        for field_key, field_mapping in self._mapping.fields.items():
            if field_mapping.excel_cell is None:
                continue
            sheet_name = field_mapping.excel_sheet or wb.sheetnames[0]
            ws = wb[sheet_name]
            cell_value = ws[field_mapping.excel_cell].value
            if cell_value is not None:
                raw[field_mapping.cir_path] = str(cell_value)

        wb.close()
        return self._build_dmp(raw)

    def _build_dmp(self, raw: dict[str, str | None]) -> DMP:
        """Build a DMP from raw cell values.

        Converts flat dot-path keyed values into the nested structure
        expected by the DMP Pydantic model, providing defaults for
        required fields not present in the source file.
        """
        nested = _unflatten(raw)
        now = datetime.now(timezone.utc)

        # Ensure required top-level fields have defaults
        nested.setdefault("title", nested.get("project", [{}])[0].get("title", "Untitled DMP"))
        default_dmp_id = {"identifier": f"import-{self._mapping.format_id}", "type": "other"}
        nested.setdefault("dmp_id", default_dmp_id)
        nested.setdefault("language", "jpn")
        nested.setdefault("ethical_issues_exist", "unknown")

        # Handle created/modified — may come from cell as string
        if "created" in nested:
            if isinstance(nested["created"], str):
                nested["created"] = _parse_datetime(nested["created"]) or now
        else:
            nested["created"] = now
        if "modified" in nested:
            if isinstance(nested["modified"], str):
                nested["modified"] = _parse_datetime(nested["modified"]) or now
        else:
            nested["modified"] = now

        # Ensure contact exists with required fields
        if "contact" not in nested:
            nested["contact"] = {}
        contact = nested["contact"]
        contact.setdefault("name", "Unknown")
        contact.setdefault("mbox", "unknown@example.com")
        contact.setdefault("contact_id", {"identifier": "unknown", "type": "other"})

        # Ensure projects have required fields
        for proj in nested.get("project", []):
            proj.setdefault("title", "Untitled Project")

        # Ensure contributors have required fields
        for contrib in nested.get("contributor", []):
            contrib.setdefault("contributor_id", {"identifier": "unknown", "type": "other"})
            contrib.setdefault("name", "Unknown")

        # Ensure at least one dataset exists
        if "dataset" not in nested:
            nested["dataset"] = [{}]
        for ds in nested["dataset"]:
            ds.setdefault("title", "Untitled Dataset")
            ds.setdefault("dataset_id", {"identifier": "unknown", "type": "other"})
            ds.setdefault("personal_data", "unknown")
            ds.setdefault("sensitive_data", "unknown")
            # Ensure distributions have required fields
            for dist in ds.get("distribution", []):
                dist.setdefault("title", ds.get("title", "Untitled"))
                dist.setdefault("data_access", "closed")
                # Ensure host has required fields
                if "host" in dist:
                    dist["host"].setdefault("title", "Unknown Repository")
                    dist["host"].setdefault("url", "")
            # Ensure security_and_privacy entries have required fields
            for sp in ds.get("security_and_privacy", []):
                sp.setdefault("title", "Security and Privacy")

        return DMP.model_validate(nested)


def _unflatten(flat: dict[str, str | None]) -> dict:
    """Convert a flat dict with dot-path keys into a nested dict.

    Numeric path segments (e.g. '0') are treated as list indices.

    Example::

        {"project.0.title": "foo"} -> {"project": [{"title": "foo"}]}
    """
    root: dict = {}
    for dotted_key, value in flat.items():
        if value is None:
            continue
        parts = dotted_key.split(".")
        _set_nested(root, parts, value)
    return root


def _set_nested(obj: dict, parts: list[str], value: object) -> None:
    """Set a value in a nested dict/list structure, creating containers as needed."""
    for i, part in enumerate(parts[:-1]):
        next_part = parts[i + 1]
        is_next_index = next_part.isdigit()

        if part.isdigit():
            idx = int(part)
            # obj is expected to be a list here
            while len(obj) <= idx:  # type: ignore[arg-type]
                obj.append({} if not is_next_index else [])  # type: ignore[union-attr]
            obj = obj[idx]  # type: ignore[index]
        else:
            if part not in obj:
                obj[part] = [] if is_next_index else {}
            obj = obj[part]

    # Set the final value
    last = parts[-1]
    if last.isdigit():
        idx = int(last)
        while len(obj) <= idx:  # type: ignore[arg-type]
            obj.append(None)  # type: ignore[union-attr]
        obj[idx] = value  # type: ignore[index]
    else:
        obj[last] = value  # type: ignore[index]


def _parse_datetime(value: str) -> datetime | None:
    """Try to parse a datetime string from various formats."""
    for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(value, fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    return None
