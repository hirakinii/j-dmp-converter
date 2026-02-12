"""Excel (XLSX) writer implementation.

Injects CIR data into agency-specific Excel templates based
on mapping definitions.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from adapters.mapping import MappingDefinition
from adapters.writers.base import BaseWriter
from core.models import DMP, MissingField, ValidationResult


class ExcelWriter(BaseWriter):
    """Writes CIR data into an Excel template file."""

    def __init__(self, mapping: MappingDefinition, template_path: Path) -> None:
        self._mapping = mapping
        self._template_path = template_path

    def write(self, dmp: DMP, output_path: Path) -> Path:
        """Inject CIR data into the Excel template and save."""
        wb = openpyxl.load_workbook(self._template_path)
        flat = self._flatten_dmp(dmp)

        for field_key, field_mapping in self._mapping.fields.items():
            if field_mapping.excel_cell is None:
                continue
            sheet_name = field_mapping.excel_sheet or wb.sheetnames[0]
            ws = wb[sheet_name]
            value = flat.get(field_mapping.cir_path)
            if value is not None:
                ws[field_mapping.excel_cell] = value

        wb.save(output_path)
        wb.close()
        return output_path

    def validate(self, dmp: DMP) -> ValidationResult:
        """Check whether the DMP has all required fields for this format."""
        flat = self._flatten_dmp(dmp)
        missing: list[MissingField] = []

        for field_key, field_mapping in self._mapping.fields.items():
            if not field_mapping.required:
                continue
            value = flat.get(field_mapping.cir_path)
            if value is None or (isinstance(value, str) and value.strip() == ""):
                missing.append(
                    MissingField(
                        field_path=field_mapping.cir_path,
                        field_label=field_key,
                        required_by=self._mapping.format_id,
                    )
                )

        return ValidationResult(is_complete=len(missing) == 0, missing_fields=missing)

    def _flatten_dmp(self, dmp: DMP) -> dict[str, str | None]:
        """Flatten the nested DMP structure into dot-path keyed values.

        E.g. 'project.title' -> 'Our Project'
        """
        flat: dict[str, str | None] = {}
        data = dmp.model_dump()
        self._flatten_dict(data, "", flat)
        return flat

    def _flatten_dict(
        self, obj: dict | list | object, prefix: str, result: dict[str, str | None]
    ) -> None:
        if isinstance(obj, dict):
            for key, value in obj.items():
                new_key = f"{prefix}.{key}" if prefix else key
                self._flatten_dict(value, new_key, result)
        elif isinstance(obj, list):
            for i, item in enumerate(obj):
                new_key = f"{prefix}.{i}"
                self._flatten_dict(item, new_key, result)
        else:
            result[prefix] = str(obj) if obj is not None else None
