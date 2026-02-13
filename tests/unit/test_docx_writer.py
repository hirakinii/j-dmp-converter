"""Tests for the DOCX writer (Step 3.1 of implementation plan).

These tests verify that CIR data is correctly rendered into
Word templates using docxtpl (Jinja2 for Word).
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from adapters.mapping import load_mapping
from adapters.writers.docx_writer import DocxWriter
from core.models import DMP

DEFINITIONS_DIR = Path(__file__).resolve().parent.parent.parent / "adapters" / "definitions"
TEMPLATES_DIR = Path(__file__).resolve().parent.parent.parent / "adapters" / "templates"


class TestCfaDocxWriter:
    """Test CFA DMP writer with Word template and Jinja2 tags."""

    @pytest.fixture
    def writer(self) -> DocxWriter:
        mapping = load_mapping(DEFINITIONS_DIR / "cfa_dmp.yaml")
        return DocxWriter(mapping=mapping, template_path=TEMPLATES_DIR / "cfa_template.docx")

    def test_write_contact_name(
        self, writer: DocxWriter, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject contact name '山田太郎' into the generated document."""
        output = tmp_path / "output.docx"
        writer.write(sample_dmp, output)

        doc = Document(output)
        all_text = "\n".join(p.text for p in doc.paragraphs)
        all_text += "\n".join(
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        )
        assert "山田太郎" in all_text

    def test_write_project_title(
        self, writer: DocxWriter, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject project title into the generated document."""
        output = tmp_path / "output.docx"
        writer.write(sample_dmp, output)

        doc = Document(output)
        all_text = "\n".join(
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        )
        assert "サンプル研究プロジェクト" in all_text

    def test_write_dataset_title(
        self, writer: DocxWriter, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject dataset title into the generated document."""
        output = tmp_path / "output.docx"
        writer.write(sample_dmp, output)

        doc = Document(output)
        all_text = "\n".join(
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        )
        assert "実験データセット" in all_text

    def test_write_affiliation(
        self, writer: DocxWriter, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject PI affiliation into the generated document."""
        output = tmp_path / "output.docx"
        writer.write(sample_dmp, output)

        doc = Document(output)
        all_text = "\n".join(
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        )
        assert "東京大学" in all_text

    def test_write_data_access(
        self, writer: DocxWriter, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject data access policy into the generated document."""
        output = tmp_path / "output.docx"
        writer.write(sample_dmp, output)

        doc = Document(output)
        all_text = "\n".join(
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        )
        assert "open" in all_text

    def test_write_repository_url(
        self, writer: DocxWriter, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject repository URL into the generated document."""
        output = tmp_path / "output.docx"
        writer.write(sample_dmp, output)

        doc = Document(output)
        all_text = "\n".join(
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        )
        assert "https://repository.example.ac.jp" in all_text

    def test_write_returns_output_path(
        self, writer: DocxWriter, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should return the output path."""
        output = tmp_path / "output.docx"
        result = writer.write(sample_dmp, output)
        assert result == output
        assert output.exists()


class TestCfaDocxWriterValidation:
    """Test validation for the CFA DMP writer."""

    def test_validate_complete_dmp(self, sample_dmp: DMP) -> None:
        """A complete DMP should pass CFA validation."""
        mapping = load_mapping(DEFINITIONS_DIR / "cfa_dmp.yaml")
        writer = DocxWriter(
            mapping=mapping, template_path=TEMPLATES_DIR / "cfa_template.docx"
        )
        result = writer.validate(sample_dmp)
        assert result.is_complete
