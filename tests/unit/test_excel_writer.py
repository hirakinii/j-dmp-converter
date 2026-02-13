"""Tests for the Excel writer (Step 1.2 of implementation plan).

These tests verify that CIR data is correctly injected into
Excel templates based on mapping definitions.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from adapters.mapping import load_mapping
from adapters.writers.excel_writer import ExcelWriter
from core.models import DMP

DEFINITIONS_DIR = Path(__file__).resolve().parent.parent.parent / "adapters" / "definitions"
TEMPLATES_DIR = Path(__file__).resolve().parent.parent.parent / "adapters" / "templates"


class TestAmedExcelWriter:
    """Test AMED DMP writer with real template and mapping coordinates."""

    def test_write_contact_name(
        self, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject contact name into cell D12."""
        mapping = load_mapping(DEFINITIONS_DIR / "amed_dmp.yaml")
        writer = ExcelWriter(mapping=mapping, template_path=TEMPLATES_DIR / "amed_template.xlsx")

        output = tmp_path / "output.xlsx"
        writer.write(sample_dmp, output)

        wb = openpyxl.load_workbook(output)
        ws = wb["DMP様式"]
        assert ws["D12"].value == "山田太郎"
        wb.close()

    def test_write_project_title(
        self, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject project title into cell D8."""
        mapping = load_mapping(DEFINITIONS_DIR / "amed_dmp.yaml")
        writer = ExcelWriter(mapping=mapping, template_path=TEMPLATES_DIR / "amed_template.xlsx")

        output = tmp_path / "output.xlsx"
        writer.write(sample_dmp, output)

        wb = openpyxl.load_workbook(output)
        ws = wb["DMP様式"]
        assert ws["D8"].value == "サンプル研究プロジェクト"
        wb.close()

    def test_write_dataset_title(
        self, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject dataset title into cell D17."""
        mapping = load_mapping(DEFINITIONS_DIR / "amed_dmp.yaml")
        writer = ExcelWriter(mapping=mapping, template_path=TEMPLATES_DIR / "amed_template.xlsx")

        output = tmp_path / "output.xlsx"
        writer.write(sample_dmp, output)

        wb = openpyxl.load_workbook(output)
        ws = wb["DMP様式"]
        assert ws["D17"].value == "実験データセット"
        wb.close()

    def test_write_affiliation(
        self, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject PI affiliation into cell D10."""
        mapping = load_mapping(DEFINITIONS_DIR / "amed_dmp.yaml")
        writer = ExcelWriter(mapping=mapping, template_path=TEMPLATES_DIR / "amed_template.xlsx")

        output = tmp_path / "output.xlsx"
        writer.write(sample_dmp, output)

        wb = openpyxl.load_workbook(output)
        ws = wb["DMP様式"]
        assert ws["D10"].value == "東京大学"
        wb.close()

    def test_write_data_access(
        self, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject data access policy into cell D25."""
        mapping = load_mapping(DEFINITIONS_DIR / "amed_dmp.yaml")
        writer = ExcelWriter(mapping=mapping, template_path=TEMPLATES_DIR / "amed_template.xlsx")

        output = tmp_path / "output.xlsx"
        writer.write(sample_dmp, output)

        wb = openpyxl.load_workbook(output)
        ws = wb["DMP様式"]
        assert ws["D25"].value == "open"
        wb.close()

    def test_write_repository_url(
        self, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject repository URL into cell D32."""
        mapping = load_mapping(DEFINITIONS_DIR / "amed_dmp.yaml")
        writer = ExcelWriter(mapping=mapping, template_path=TEMPLATES_DIR / "amed_template.xlsx")

        output = tmp_path / "output.xlsx"
        writer.write(sample_dmp, output)

        wb = openpyxl.load_workbook(output)
        ws = wb["DMP様式"]
        assert ws["D32"].value == "https://repository.example.ac.jp"
        wb.close()


class TestJspsExcelWriter:
    """Test JSPS DMP writer with real template and mapping coordinates."""

    def test_write_contact_name(
        self, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject contact name into cell D13."""
        mapping = load_mapping(DEFINITIONS_DIR / "jsps_dmp.yaml")
        writer = ExcelWriter(mapping=mapping, template_path=TEMPLATES_DIR / "jsps_template.xlsx")

        output = tmp_path / "output.xlsx"
        writer.write(sample_dmp, output)

        wb = openpyxl.load_workbook(output)
        ws = wb["DMP様式例"]
        assert ws["D13"].value == "山田太郎"
        wb.close()

    def test_write_dataset_title(
        self, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject dataset title into cell B22."""
        mapping = load_mapping(DEFINITIONS_DIR / "jsps_dmp.yaml")
        writer = ExcelWriter(mapping=mapping, template_path=TEMPLATES_DIR / "jsps_template.xlsx")

        output = tmp_path / "output.xlsx"
        writer.write(sample_dmp, output)

        wb = openpyxl.load_workbook(output)
        ws = wb["DMP様式例"]
        assert ws["B22"].value == "実験データセット"
        wb.close()

    def test_write_erad_project_id(
        self, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject e-Rad project ID into cell C9."""
        mapping = load_mapping(DEFINITIONS_DIR / "jsps_dmp.yaml")
        writer = ExcelWriter(mapping=mapping, template_path=TEMPLATES_DIR / "jsps_template.xlsx")

        output = tmp_path / "output.xlsx"
        writer.write(sample_dmp, output)

        wb = openpyxl.load_workbook(output)
        ws = wb["DMP様式例"]
        assert ws["C9"].value == "JP26000001"
        wb.close()


class TestMetiExcelWriter:
    """Test METI DMP writer with real template and mapping coordinates."""

    def test_write_project_title(
        self, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject project title into cell D5."""
        mapping = load_mapping(DEFINITIONS_DIR / "meti_dmp.yaml")
        writer = ExcelWriter(mapping=mapping, template_path=TEMPLATES_DIR / "meti_template.xlsx")

        output = tmp_path / "output.xlsx"
        writer.write(sample_dmp, output)

        wb = openpyxl.load_workbook(output)
        ws = wb["様式"]
        assert ws["D5"].value == "サンプル研究プロジェクト"
        wb.close()

    def test_write_dataset_title(
        self, sample_dmp: DMP, tmp_path: Path
    ) -> None:
        """Writer should inject dataset title into cell C9."""
        mapping = load_mapping(DEFINITIONS_DIR / "meti_dmp.yaml")
        writer = ExcelWriter(mapping=mapping, template_path=TEMPLATES_DIR / "meti_template.xlsx")

        output = tmp_path / "output.xlsx"
        writer.write(sample_dmp, output)

        wb = openpyxl.load_workbook(output)
        ws = wb["様式"]
        assert ws["C9"].value == "実験データセット"
        wb.close()


class TestExcelWriterValidation:
    def test_validate_complete_dmp_for_amed(self, sample_dmp: DMP) -> None:
        """A complete DMP should pass AMED validation."""
        mapping = load_mapping(DEFINITIONS_DIR / "amed_dmp.yaml")
        writer = ExcelWriter(
            mapping=mapping, template_path=TEMPLATES_DIR / "amed_template.xlsx"
        )
        result = writer.validate(sample_dmp)
        assert result.is_complete

    def test_validate_detects_missing_jsps_fields(self, sample_dmp: DMP) -> None:
        """JSPS requires more fields — validation should detect gaps."""
        mapping = load_mapping(DEFINITIONS_DIR / "jsps_dmp.yaml")
        writer = ExcelWriter(
            mapping=mapping, template_path=TEMPLATES_DIR / "jsps_template.xlsx"
        )
        result = writer.validate(sample_dmp)
        # JSPS requires data_access_details (distribution description)
        # which the sample_dmp does not provide
        missing_paths = [f.field_path for f in result.missing_fields]
        assert "dataset.0.distribution.0.description" in missing_paths
