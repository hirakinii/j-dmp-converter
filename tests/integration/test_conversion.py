"""Integration tests for the full conversion pipeline.

Tests the complete Read -> CIR -> Write flow across different
agency formats, verifying that data is preserved during conversion.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from adapters.mapping import load_mapping
from adapters.readers.excel_reader import ExcelReader
from adapters.writers.excel_writer import ExcelWriter
from core.converter import ConversionService
from core.models import DMP

DEFINITIONS_DIR = Path(__file__).resolve().parent.parent.parent / "adapters" / "definitions"
TEMPLATES_DIR = Path(__file__).resolve().parent.parent.parent / "adapters" / "templates"


def _build_service() -> ConversionService:
    """Create a ConversionService with all readers and writers registered."""
    service = ConversionService()

    for format_id, definition_file, template_file in [
        ("amed", "amed_dmp.yaml", "amed_template.xlsx"),
        ("jsps", "jsps_dmp.yaml", "jsps_template.xlsx"),
        ("meti", "meti_dmp.yaml", "meti_template.xlsx"),
    ]:
        mapping = load_mapping(DEFINITIONS_DIR / definition_file)
        service.register_reader(format_id, ExcelReader(mapping=mapping))
        service.register_writer(
            format_id,
            ExcelWriter(mapping=mapping, template_path=TEMPLATES_DIR / template_file),
        )

    return service


class TestConversionPipeline:
    """Test the full conversion pipeline between formats."""

    def test_amed_write_then_read(self, sample_dmp: DMP, tmp_path: Path) -> None:
        """Write CIR to AMED format, then read it back."""
        service = _build_service()

        output = tmp_path / "amed_output.xlsx"
        service.write(sample_dmp, "amed", output)
        recovered = service.read("amed", output)

        assert recovered.contact.name == sample_dmp.contact.name
        assert recovered.project[0].title == sample_dmp.project[0].title
        assert recovered.dataset[0].title == sample_dmp.dataset[0].title

    def test_amed_to_jsps_conversion(self, sample_dmp: DMP, tmp_path: Path) -> None:
        """Write AMED, read it, then write as JSPS — key fields should transfer."""
        service = _build_service()

        # Write as AMED
        amed_path = tmp_path / "amed.xlsx"
        service.write(sample_dmp, "amed", amed_path)

        # Read AMED
        cir = service.read("amed", amed_path)

        # Validate against JSPS requirements
        result = service.validate(cir, "jsps")
        # Not all fields will transfer (JSPS may require more), but core data should exist
        assert cir.contact.name == "山田太郎"
        assert cir.dataset[0].title == "実験データセット"

        # Write as JSPS
        jsps_path = tmp_path / "jsps.xlsx"
        service.write(cir, "jsps", jsps_path)

        # Verify JSPS output has the transferred data
        wb = openpyxl.load_workbook(jsps_path)
        ws = wb["DMP様式例"]
        assert ws["D13"].value == "山田太郎"
        assert ws["B22"].value == "実験データセット"
        wb.close()

    def test_amed_to_meti_conversion(self, sample_dmp: DMP, tmp_path: Path) -> None:
        """Write AMED, read it, then write as METI — shared fields should transfer."""
        service = _build_service()

        # Write as AMED
        amed_path = tmp_path / "amed.xlsx"
        service.write(sample_dmp, "amed", amed_path)

        # Read AMED -> CIR
        cir = service.read("amed", amed_path)

        # Write as METI
        meti_path = tmp_path / "meti.xlsx"
        service.write(cir, "meti", meti_path)

        # Verify METI output has transferred fields
        wb = openpyxl.load_workbook(meti_path)
        ws = wb["様式"]
        assert ws["D5"].value == "サンプル研究プロジェクト"
        assert ws["C9"].value == "実験データセット"
        wb.close()

    def test_convert_with_validation(self, sample_dmp: DMP, tmp_path: Path) -> None:
        """Full convert pipeline should validate and write when complete."""
        service = _build_service()

        # Write AMED first
        amed_path = tmp_path / "amed.xlsx"
        service.write(sample_dmp, "amed", amed_path)

        # Convert AMED -> AMED (same format, should be complete)
        output, result = service.convert("amed", "amed", amed_path, tmp_path / "out.xlsx")
        assert result.is_complete

    def test_available_formats(self) -> None:
        """Service should list all registered formats."""
        service = _build_service()
        assert sorted(service.available_readers()) == ["amed", "jsps", "meti"]
        assert sorted(service.available_writers()) == ["amed", "jsps", "meti"]
